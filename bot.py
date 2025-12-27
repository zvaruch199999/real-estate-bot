import asyncio, os
from aiogram import Bot, Dispatcher, F
from aiogram.types import Message, InlineKeyboardButton, InlineKeyboardMarkup, InputMediaPhoto
from aiogram.fsm.state import StatesGroup, State
from aiogram.fsm.context import FSMContext
from aiogram.fsm.storage.memory import MemoryStorage
from openpyxl import Workbook, load_workbook

TOKEN = os.getenv("BOT_TOKEN")
GROUP_ID = int(os.getenv("GROUP_CHAT_ID"))

bot = Bot(TOKEN)
dp = Dispatcher(storage=MemoryStorage())

# ===================== EXCEL =====================
FILE = "deals.xlsx"

def save_deal(row):
    if not os.path.exists(FILE):
        wb = Workbook()
        ws = wb.active
        ws.append([
            "ID","Категорія","Тип","Адреса","Ціна","Маклер",
            "Хто знайшов житло","Хто клієнта","Дата",
            "Комісія","Оплати","Графік","ПІБ","Контакт"
        ])
        wb.save(FILE)

    wb = load_workbook(FILE)
    ws = wb.active
    ws.append(row)
    wb.save(FILE)

# ===================== STATES =====================
class Offer(StatesGroup):
    category = State()
    type = State()
    address = State()
    price = State()
    broker = State()
    photos = State()
    confirm = State()

class CloseDeal(StatesGroup):
    step = State()

# ===================== CREATE =====================
@dp.message(F.text == "/start")
async def start(m: Message, s: FSMContext):
    await s.clear()
    await m.answer("Напишіть `create` щоб створити пропозицію")

@dp.message(F.text.lower() == "create")
async def create(m: Message, s: FSMContext):
    await s.set_state(Offer.category)
    await m.answer("Категорія:")

@dp.message(Offer.category)
async def cat(m: Message, s: FSMContext):
    await s.update_data(category=m.text)
    await s.set_state(Offer.type)
    await m.answer("Тип:")

@dp.message(Offer.type)
async def typ(m: Message, s: FSMContext):
    await s.update_data(type=m.text)
    await s.set_state(Offer.address)
    await m.answer("Адреса:")

@dp.message(Offer.address)
async def addr(m: Message, s: FSMContext):
    await s.update_data(address=m.text)
    await s.set_state(Offer.price)
    await m.answer("Ціна:")

@dp.message(Offer.price)
async def price(m: Message, s: FSMContext):
    await s.update_data(price=m.text)
    await s.set_state(Offer.broker)
    await m.answer("Маклер (@username):")

@dp.message(Offer.broker)
async def broker(m: Message, s: FSMContext):
    await s.update_data(broker=m.text, photos=[])
    await s.set_state(Offer.photos)
    await m.answer("Надішліть фото. Коли завершите — `/done`")

@dp.message(F.photo, Offer.photos)
async def photos(m: Message, s: FSMContext):
    data = await s.get_data()
    data["photos"].append(m.photo[-1].file_id)
    await s.update_data(photos=data["photos"])
    await m.answer(f"Фото додано ({len(data['photos'])})")

@dp.message(F.text == "/done", Offer.photos)
async def done(m: Message, s: FSMContext):
    d = await s.get_data()
    text = (
        f"🏠 НОВА ПРОПОЗИЦІЯ\n\n"
        f"📦 Категорія: {d['category']}\n"
        f"🏡 Тип: {d['type']}\n"
        f"📍 Адреса: {d['address']}\n"
        f"💰 Ціна: {d['price']}\n"
        f"👤 Маклер: {d['broker']}\n\n"
        "Напишіть:\n"
        "✅ publish — опублікувати\n"
        "✏️ edit — змінити\n"
        "❌ cancel — скасувати"
    )
    await s.set_state(Offer.confirm)
    await m.answer(text)

# ===================== PUBLISH =====================
@dp.message(F.text == "publish", Offer.confirm)
async def publish(m: Message, s: FSMContext):
    d = await s.get_data()

    media = [InputMediaPhoto(media=p) for p in d["photos"]]
    media[0].caption = (
        f"🏠 НОВА ПРОПОЗИЦІЯ\n"
        f"🟢 Статус: Актуально\n\n"
        f"📦 {d['category']}\n"
        f"🏡 {d['type']}\n"
        f"📍 {d['address']}\n"
        f"💰 {d['price']}\n"
        f"👤 {d['broker']}"
    )

    kb = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="🟢 Актуально", callback_data="status_active"),
        InlineKeyboardButton(text="🟡 Резерв", callback_data="status_reserved"),
        InlineKeyboardButton(text="🔴 Неактуально", callback_data="status_closed"),
        InlineKeyboardButton(text="🔒 Закрити угоду", callback_data="close")
    ]])

    msgs = await bot.send_media_group(GROUP_ID, media)
    await bot.send_message(GROUP_ID, "⬆️ Керування статусом", reply_markup=kb)

    await m.answer("✅ Пропозицію опубліковано")
    await s.clear()

# ===================== CLOSE DEAL =====================
@dp.callback_query(F.data == "close")
async def close(cb, s: FSMContext):
    await s.set_state(CloseDeal.step)
    await s.update_data(step=0, answers=[])
    await cb.message.answer("Хто знайшов нерухомість?")

@dp.message(CloseDeal.step)
async def close_steps(m: Message, s: FSMContext):
    data = await s.get_data()
    answers = data["answers"]
    answers.append(m.text)

    questions = [
        "Хто знайшов клієнта?",
        "Дата контракту:",
        "Сума комісії:",
        "Кількість оплат:",
        "Графік оплат:",
        "ПІБ клієнта:",
        "Контакт клієнта:"
    ]

    if len(answers) < len(questions):
        await s.update_data(answers=answers)
        await m.answer(questions[len(answers)-1])
    else:
        save_deal(["#", "", "", "", "", "", *answers])
        await m.answer("✅ Угоду закрито та збережено в Excel")
        await s.clear()

# ===================== RUN =====================
async def main():
    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())
